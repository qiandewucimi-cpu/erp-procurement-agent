from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


BASE_DIR = Path(__file__).resolve().parent


def write_sample(filename: str, rows: list[list[object]]) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "BOM"
    sheet.append(["物料编码", "物料名称", "数量", "单位"])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for width, column in zip((18, 22, 12, 10), ("A", "B", "C", "D")):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(BASE_DIR / "samples" / filename)


def write_material_template(filename: str, rows: list[list[object]]) -> None:
    """生成物料档案模板：用户照此填写自己的物料主数据，再导入模拟 ERP。"""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "物料档案"
    sheet.append(["物料编码", "物料名称", "供应商编码", "供应商名称", "单价", "包装费", "币种"])
    for row in rows:
        sheet.append(row)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for width, column in zip((16, 22, 14, 20, 10, 10, 8), ("A", "B", "C", "D", "E", "F", "G")):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    workbook.save(BASE_DIR / "samples" / filename)


if __name__ == "__main__":
    write_sample(
        "正常示例_BOM.xlsx",
        [
            ["MAT-FAB-001", "再生涤纶面料", 1200, "米"],
            ["MAT-ZIP-002", "尼龙拉链", 800, "条"],
            ["MAT-LBL-003", "洗水唛", 800, "个"],
        ],
    )
    write_sample(
        "异常示例_BOM.xlsx",
        [
            ["MAT-FAB-001", "涤纶面料（旧名称）", 1200, "米"],
            ["MAT-UNKNOWN-999", "未建档辅料", 500, "个"],
            ["MAT-ZIP-002", "尼龙拉链", 0, "条"],
        ],
    )
    write_material_template(
        "物料档案模板.xlsx",
        [
            ["MY-FAB-100", "自定义面料（示例）", "SUP-900", "自定义供应商（示例）", 25.80, 0.40, "CNY"],
            ["MY-ACC-200", "自定义辅料（示例）", "SUP-900", "自定义供应商（示例）", 3.20, 0.10, "CNY"],
        ],
    )
    print("已生成 2 份完全合成的 BOM 示例与 1 份物料档案模板。")

